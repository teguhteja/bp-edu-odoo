"""Konversi panduan Markdown di docs/labapp ke .docx (python-docx), gambar ikut tertanam.

Mendukung subset Markdown yang dipakai panduan: heading, paragraf, list (- / 1.), tabel,
gambar, **tebal**, `kode`, [tautan](url), dan garis pemisah.

    python scripts/labapp_docs/md_to_docx.py            # semua panduan-*.md
    python scripts/labapp_docs/md_to_docx.py docs/labapp/panduan-dosen.md
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[2] / 'docs' / 'labapp'
INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)')
IMG = re.compile(r'^!\[([^\]]*)\]\(([^)]+)\)$')
LIST = re.compile(r'^(\s*)([-*]|\d+\.)\s+(.*)$')


def tambah_inline(par, teks):
    for bagian in INLINE.split(teks):
        if not bagian:
            continue
        if bagian.startswith('**') and bagian.endswith('**'):
            par.add_run(bagian[2:-2]).bold = True
        elif bagian.startswith('`') and bagian.endswith('`'):
            run = par.add_run(bagian[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif bagian.startswith('['):
            par.add_run(re.match(r'\[([^\]]+)\]', bagian).group(1))
        elif bagian.startswith('*') and bagian.endswith('*'):
            par.add_run(bagian[1:-1]).italic = True
        else:
            par.add_run(bagian)


def tambah_gambar(doc, path, alt, base):
    file = (base / path).resolve()
    if not file.exists():
        raise FileNotFoundError(file)
    doc.add_picture(str(file), width=Cm(15.5))
    if alt:
        cap = doc.add_paragraph()
        run = cap.add_run(alt)
        run.italic = True
        run.font.size = Pt(9)


def tambah_tabel(doc, baris):
    rows = [[c.strip() for c in b.strip().strip('|').split('|')] for b in baris
            if not re.match(r'^\|?\s*:?-{3,}', b.strip())]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, teks in enumerate(row[:len(rows[0])]):
            cell = table.cell(i, j)
            cell.text = ''
            tambah_inline(cell.paragraphs[0], teks)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def konversi(md_path):
    md_path = Path(md_path)
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)

    lines = md_path.read_text(encoding='utf-8').splitlines()
    para = []
    jumlah_gambar = 0

    def flush():
        if para:
            tambah_inline(doc.add_paragraph(), ' '.join(s.strip() for s in para))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
        elif stripped.startswith('#'):
            flush()
            level = len(stripped) - len(stripped.lstrip('#'))
            doc.add_heading(stripped[level:].strip(), level=min(level - 1, 3) if level > 1 else 0)
        elif stripped == '---':
            flush()
        elif IMG.match(stripped):
            flush()
            alt, path = IMG.match(stripped).groups()
            tambah_gambar(doc, path, alt, md_path.parent)
            jumlah_gambar += 1
        elif stripped.startswith('|'):
            flush()
            blok = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                blok.append(lines[i])
                i += 1
            tambah_tabel(doc, blok)
            continue
        elif LIST.match(line):
            flush()
            indent, penanda, teks = LIST.match(line).groups()
            # gabungkan baris lanjutan item list (indentasi, bukan item/gambar/tabel baru)
            while i + 1 < len(lines) and lines[i + 1].startswith(' ') and lines[i + 1].strip() \
                    and not LIST.match(lines[i + 1]) and not IMG.match(lines[i + 1].strip()) \
                    and not lines[i + 1].strip().startswith('|'):
                i += 1
                teks += ' ' + lines[i].strip()
            if penanda[0].isdigit():
                # nomor ditulis apa adanya: style "List Number" Word tidak mengulang dari 1 per daftar
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 + 0.6 * (len(indent) // 2))
                p.paragraph_format.first_line_indent = Cm(-0.6)
                tambah_inline(p, f'{penanda} {teks}')
            else:
                p = doc.add_paragraph(style='List Bullet 2' if len(indent) >= 2 else 'List Bullet')
                tambah_inline(p, teks)
        else:
            para.append(line)
        i += 1
    flush()

    out = md_path.with_suffix('.docx')
    doc.save(out)
    print(f'{out.name}: {jumlah_gambar} gambar')
    return out


if __name__ == '__main__':
    targets = sys.argv[1:] or sorted(DOCS.glob('panduan-*.md'))
    for target in targets:
        konversi(target)
